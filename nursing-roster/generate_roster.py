"""Generate the nursing roster for Oct-Dec 2026 from the September 2026 roster.

The September sheet follows fixed rules, which this script infers per nurse
and then continues day by day:
  * every nurse has one fixed weekly off day (O);
  * after each off day the shift rotates A -> C -> B -> A;
  * GENERAL duty (G) and the fixed-morning nurse (Swati) do not rotate.

Run:  python3 generate_roster.py
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

HERE = Path(__file__).resolve().parent
OUT_XLSX = HERE / "Nursing_Roster_Oct-Dec_2026.xlsx"
OUT_ICU1ST_XLSX = HERE / "ICU1ST_Roster_Oct-Dec_2026.xlsx"

SEPT_START = dt.date(2026, 9, 1)
MONTHS = [(2026, 10), (2026, 11), (2026, 12)]
NEXT_SHIFT = {"A": "C", "C": "B", "B": "A"}
WEEKDAY = ["M", "T", "W", "TH", "F", "S", "SU"]

# September 2026 roster, transcribed from the printed sheet (1st-30th).
# "R" marks days after a resignation.
SEPTEMBER = [
    ("ICU1ST", [
        ("SOURABH",     "BBBBBOAAAAAAOCCCCCCOBBBBBBOAAA"),
        ("SHIV P.",     "AAAAAAOCCCCCCOBBBBBBOAAAAAAOCC"),
        ("ASHUTOSH",    "OCCCCCCOBBBBBBOAAAAAAOCCCCCCOB"),
        ("AJEET",       "COBBBBBBOAAAAAAOCCCCCCOBBBBBBO"),
        ("SUMAN",       "COBBBBBBOAAAAAAOCCCCCCOBBBBBBO"),
        ("MANSI",       "OCCCCCCOBBBBBBOAAAAAAOCCCCCCOB"),
        ("MOHIT",       "BBOAAAAAAOCCCCCCOBBBBBBOAAAAAA"),
        ("ABIGEL",      "GGGGGOGGGGGGOGGGGGGOGGGGGGOGGG"),
    ]),
    ("ICU2ND", [
        ("SHARAD",      "CCCCOBBBBBBOAAAAAAOCCCCCCOBBBB"),
        ("RAVINDRA",    "BBOAAAAAAOCCCCCCOBBBBBBOAAAAAA"),
        ("SOHEL",       "BBBOAAAAAAOCCCCCCOBBBBBBOAAAAA"),
        ("DINESH",      "AAAOCCCCCCOBBBBBBOAAAAAAOCCCCC"),
        ("VANDANA R.",  "CCCCOBBBBBBOAAAAAAOCCCCCCOBBBB"),
        ("SUDHA",       "BBBOAAAAAAOCCCCCCOBBBBBBOAAAAA"),
        ("DURGESHWARI", "AAOCCCCCCOBBBBBBOAAAAAAOCCCCCC"),
        ("SWATI",       "AOAAAAAAOAAAAAAOAAAAAAOAAAAAAA"),
        ("DEVKI",       "CCOBBBBBBOAAAAAAOCCCCCCOBBBBBB"),
        ("VARSHA",      "AAAOCCCCCCOBBBBBBOAAAAAAOCCCCC"),
        ("SANGEETA G.", "AAAAAAOCCCCCCOAAAARRRRRRRRRRRR"),
    ]),
    ("GENERAL", [
        ("POONAM S.",   "CCCCCCOBBBBBBOAAAAAAOCCCCCCOBB"),
        ("RITA",        "OBBBBBBOAAAAAAOCCCCCCOBBBBBBOA"),
        ("NEETU",       "BOAAAAAAOCCCCCCOBBBBBBOAAAAAAO"),
        ("SONU",        "OBBBBBBOAAAAAAOCCCCCCOBBBBBBOA"),
        ("POONAM N.",   "CCCCOBBBBBBOAAAAAAOCCCCCCOBBBB"),
        ("VANDANA V.",  "AAAAOCCCCCCOBBBBBBOAAAAAAOCCCC"),
    ]),
    ("PVT", [
        ("POONAM A.",   "BBBBBOAAAAAAOCCCCCCOBBBBBBOAAA"),
    ]),
]


def simulate(first_block, off_weekday, rotates, start, days):
    """Duty codes for `days` days from `start` under the roster rules."""
    shift, out = first_block, []
    for i in range(days):
        day = start + dt.timedelta(days=i)
        if day.weekday() == off_weekday:
            out.append("O")
            if rotates:
                shift = NEXT_SHIFT[shift]
        else:
            out.append(shift)
    return out


def infer_rule(name, sept):
    """Find the weekly off day and starting shift that best reproduce September."""
    worked = sept.replace("R", "")
    offs = {(SEPT_START + dt.timedelta(days=i)).weekday()
            for i, c in enumerate(worked) if c == "O"}
    if len(offs) != 1:
        raise ValueError(f"{name}: off days fall on more than one weekday")
    off = offs.pop()
    rotates = len(set(worked) - {"O"}) > 1
    best = None
    for first in ("ABC" if rotates else sorted(set(worked) - {"O"})):
        sim = simulate(first, off, rotates, SEPT_START, len(worked))
        diffs = [i + 1 for i, (a, b) in enumerate(zip(sim, worked)) if a != b]
        if best is None or len(diffs) < len(best[2]):
            best = (first, rotates, diffs)
    first, rotates, diffs = best
    return {"off": off, "first": first, "rotates": rotates, "sept_diffs": diffs}


def build_rosters():
    """Return {(year, month): [(section, [(name, codes)])]} plus inference notes."""
    end = dt.date(2026, 12, 31)
    total_days = (end - SEPT_START).days + 1
    rosters, notes = {m: [] for m in MONTHS}, []
    for section, staff in SEPTEMBER:
        rows = {m: [] for m in MONTHS}
        for name, sept in staff:
            if "R" in sept:
                notes.append(f"{name}: resigned in September - not rostered")
                continue
            rule = infer_rule(name, sept)
            if rule["sept_diffs"]:
                notes.append(f"{name}: September sheet differs from the weekly "
                             f"pattern on day(s) {rule['sept_diffs']}")
            codes = simulate(rule["first"], rule["off"], rule["rotates"],
                             SEPT_START, total_days)
            for (y, m) in MONTHS:
                first = (dt.date(y, m, 1) - SEPT_START).days
                n = (dt.date(y + m // 12, m % 12 + 1, 1) - dt.date(y, m, 1)).days
                rows[(y, m)].append((name, codes[first:first + n]))
        for m in MONTHS:
            rosters[m].append((section, rows[m]))
    return rosters, notes


# ---------------------------------------------------------------- workbook ---

FONT = "Arial"
BLACK = PatternFill("solid", fgColor="000000")
GREY = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", indent=1)


def cell(ws, row, col, value, bold=False, color="000000", fill=None,
         align=CENTER, size=10, border=True):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, bold=bold, color=color, size=size)
    c.alignment = align
    if fill:
        c.fill = fill
    if border:
        c.border = BOX
    return c


def header_row(ws, row, label, sno, days, year, month):
    """Black bar with section name, day numbers and weekday initials."""
    cell(ws, row, 1, sno, bold=True, color="FFFFFF", fill=BLACK)
    cell(ws, row, 2, label, bold=True, color="FFFFFF", fill=BLACK, align=LEFT)
    for d in range(1, days + 1):
        wd = WEEKDAY[dt.date(year, month, d).weekday()]
        cell(ws, row, d + 2, f"{d}\n{wd}", bold=True, color="FFFFFF",
             fill=BLACK, align=Alignment(horizontal="center",
                                         vertical="center", wrap_text=True),
             size=9)
    ws.row_dimensions[row].height = 26


def month_days(year, month):
    return (dt.date(year + month // 12, month % 12 + 1, 1)
            - dt.date(year, month, 1)).days


def write_title(ws, row, text, last_col):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell(ws, row, 1, text, bold=True, size=13, border=False)
    ws.row_dimensions[row].height = 22


def write_staff(ws, row, year, month, sections, first_label=None):
    """Header bar plus one line per nurse for each section; returns next row."""
    days = month_days(year, month)
    sno, section_ranges = 1, []
    for i, (section, staff) in enumerate(sections):
        label = first_label if (i == 0 and first_label) else section
        header_row(ws, row, label, "SNO" if i == 0 else "@", days, year, month)
        row += 1
        first = row
        for name, codes in staff:
            cell(ws, row, 1, sno)
            cell(ws, row, 2, name, align=LEFT)
            for d, code in enumerate(codes, start=3):
                if code == "O":
                    cell(ws, row, d, code, bold=True, color="FFFFFF", fill=BLACK)
                else:
                    cell(ws, row, d, code)
            ws.row_dimensions[row].height = 17
            row += 1
            sno += 1
        section_ranges.append((section, first, row - 1))
    return row, section_ranges


def write_legend(ws, row, last_col):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell(ws, row, 1, "A = Morning     B = Evening     C = Night     "
                     "G = General     O = Weekly Off", bold=True, border=False,
         align=LEFT)


def setup_page(ws, days):
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 17
    for d in range(1, days + 1):
        ws.column_dimensions[get_column_letter(d + 2)].width = 3.9
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4)
    ws.print_options.horizontalCentered = True


def write_month(wb, year, month, sections):
    days = month_days(year, month)
    title = dt.date(year, month, 1).strftime("%B").upper()
    ws = wb.create_sheet(title[:3])
    write_title(ws, 1, f"NURSING ROSTER MONTH OF {title} {year}", days + 2)
    row, section_ranges = write_staff(ws, 2, year, month, sections)
    write_legend(ws, row + 1, days + 2)
    # Daily head-count per section, kept live with COUNTIF so manual swaps
    # made later in Excel are reflected automatically.
    write_counts(ws, row + 3, days, section_ranges)
    ws.freeze_panes = "C3"
    setup_page(ws, days)


def write_quarter(wb, rosters, section):
    """All months stacked on one printable sheet for a single section."""
    (y1, m1), (y2, m2) = MONTHS[0], MONTHS[-1]
    first, last = dt.date(y1, m1, 1), dt.date(y2, m2, 1)
    ws = wb.create_sheet(f"{first:%b}-{last:%b}".upper())
    last_col = max(month_days(y, m) for y, m in MONTHS) + 2
    write_title(ws, 1, f"NURSING ROSTER {section}  -  {first:%B} TO "
                       f"{last:%B} {y2}".upper(), last_col)
    row = 2
    for (y, m) in MONTHS:
        staff = [(s, st) for s, st in rosters[(y, m)] if s == section]
        row, _ = write_staff(ws, row, y, m, staff,
                             first_label=dt.date(y, m, 1).strftime("%B %Y").upper())
        row += 1
    write_legend(ws, row, last_col)
    setup_page(ws, last_col - 2)


def write_counts(ws, row, days, section_ranges):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    cell(ws, row, 1, "STAFF ON DUTY (auto count)", bold=True, fill=GREY,
         align=LEFT, size=9)
    cell(ws, row, 2, None, fill=GREY)
    for d in range(1, days + 1):
        cell(ws, row, d + 2, d, bold=True, fill=GREY, size=9)
    row += 1
    for section, r1, r2 in section_ranges:
        shifts = "G" if section == "ICU1ST" else ""
        for shift in "ABC" + shifts:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            cell(ws, row, 1, f"{section} - {shift}", align=LEFT, size=9)
            cell(ws, row, 2, None)
            for d in range(1, days + 1):
                col = get_column_letter(d + 2)
                cell(ws, row, d + 2, f'=COUNTIF({col}{r1}:{col}{r2},"{shift}")',
                     size=9)
            row += 1


def save_workbook(path, rosters):
    """Write one sheet per month with every section."""
    wb = Workbook()
    wb.remove(wb.active)
    for (y, m) in MONTHS:
        write_month(wb, y, m, rosters[(y, m)])
    wb.save(path)
    print(f"Saved {path.name}")


def main():
    rosters, notes = build_rosters()
    save_workbook(OUT_XLSX, rosters)
    wb = Workbook()
    wb.remove(wb.active)
    write_quarter(wb, rosters, "ICU1ST")
    wb.save(OUT_ICU1ST_XLSX)
    print(f"Saved {OUT_ICU1ST_XLSX.name}")
    for n in notes:
        print("NOTE:", n)


if __name__ == "__main__":
    main()
